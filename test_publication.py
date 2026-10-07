import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from openpyxl import load_workbook
from src.tracker import update_tracker

class PublicationTests(unittest.TestCase):
    def test_remote_changes_preserved_and_replay_deduplicated(self):
        source = Path(__file__).resolve().parent.parent
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp); remote = base / "remote.git"; seed = base / "seed"; runner = base / "runner"
            def git(*args, cwd=base):
                return subprocess.run(["git", *args], cwd=cwd, check=True, capture_output=True, text=True)
            git("init", "--bare", str(remote))
            git("init", "-b", "main", str(seed))
            shutil.copytree(source / "src", seed / "src", ignore=shutil.ignore_patterns("__pycache__"))
            (seed / "config").mkdir()
            (seed / "config/cv_profile.yaml").write_text("main_min_score: 0\nswiss_min_score: 0\n")
            (seed / "README.md").write_text("Initial version")
            git("config", "user.name", "Fixture", cwd=seed); git("config", "user.email", "fixture@example.invalid", cwd=seed)
            git("add", ".", cwd=seed); git("commit", "-m", "Initial", cwd=seed)
            git("remote", "add", "origin", str(remote), cwd=seed); git("push", "-u", "origin", "main", cwd=seed)
            git("clone", "-b", "main", str(remote), str(runner))
            # A second writer updates both source and workbook after runner checkout.
            latest_job = {"url": "https://example.invalid/job/latest", "title": "Remote writer job", "relevance_score": 5}
            update_tracker(seed / "data/job_tracker.xlsx", [latest_job])
            (seed / "README.md").write_text("New remote edit must survive")
            git("add", ".", cwd=seed); git("commit", "-m", "Concurrent remote edit", cwd=seed); git("push", cwd=seed)
            results = base / "results.json"
            new_job = {"url": "https://example.invalid/job/scraped", "title": "Scraped job", "relevance_score": 6}
            results.write_text(json.dumps({"main_jobs": [new_job], "swiss_jobs": []}))
            for _ in range(2):
                subprocess.run(["python", "-m", "src.publish_results", str(results), "--branch", "main"], cwd=runner, check=True, capture_output=True, text=True)
            git("pull", "--ff-only", cwd=seed)
            self.assertEqual((seed / "README.md").read_text(), "New remote edit must survive")
            wb = load_workbook(seed / "data/job_tracker.xlsx")
            urls = [wb["Jobs"].cell(row=i, column=7).value for i in range(2, wb["Jobs"].max_row+1)]
            self.assertCountEqual(urls, [latest_job["url"], new_job["url"]])
            self.assertTrue((seed / "docs/index.html").exists())
