"""Consolidated website implementation. See INSTALL.txt for migration details."""
from __future__ import annotations

# ========================================================================
# generate_html: consolidated implementation
# ========================================================================



import datetime as dt
from pathlib import Path

import yaml
from urllib.parse import urlparse
from src.matching import area_priority, sort_key
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.matching import allowed_job
from src.matching import excluded_seniority
from src.tracker import COLUMNS, NEW_ROW_FILL

ROOT = Path(__file__).resolve().parent.parent
TRACKER_FILE = ROOT / "data" / "job_tracker.xlsx"
OUTPUT_FILE = ROOT / "docs" / "index.html"
CV_PROFILE_FILE = ROOT / "config" / "cv_profile.yaml"


def _rgb_matches(cell, fill) -> bool:
    cell_rgb = getattr(cell.fill.start_color, "rgb", None)
    fill_rgb = getattr(fill.start_color, "rgb", None)
    if not isinstance(cell_rgb, str) or not isinstance(fill_rgb, str):
        return False
    return cell_rgb[-6:] == fill_rgb[-6:]  # compare ignoring alpha prefix


def _escape(value) -> str:
    if value is None:
        return ""
    text = str(value)
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#x27;")
    )


def _render_table(
    ws: Worksheet | None, columns: list[str], new_row_fill, empty_message: str, row_class: str,
) -> tuple[str, int]:
    if ws is None or ws.max_row < 2:
        return f"<tr><td colspan='{len([c for c in columns if c != 'URL']) + 1}'>{_escape(empty_message)}</td></tr>", 0

    url_col = columns.index("URL") + 1
    title_col = columns.index("Job Title")

    rows_html_parts = []
    ordered_rows = sorted(range(2, ws.max_row + 1), key=lambda row: sort_key(ws.cell(row, columns.index('Location') + 1).value, ws.cell(row, columns.index('Relevance Score (1-10)') + 1).value))
    for row in ordered_rows:
        values = [ws.cell(row=row, column=c).value for c in range(1, len(columns) + 1)]
        if all(v is None for v in values):
            continue  # skip fully-empty rows
        record = dict(zip(columns, values))
        if excluded_seniority(record.get('Job Title')) or not allowed_job({'city':record.get('City'), 'company':record.get('Company'), 'url':record.get('URL'), 'location':record.get('Location')}):
            continue
        is_highlighted = _rgb_matches(ws.cell(row=row, column=1), new_row_fill)
        url = str(values[url_col - 1] or "#")
        if urlparse(url).scheme not in ("http", "https"):
            url = "#"

        rank, area = area_priority(record.get('Location'))
        cells = [f"<td>{_escape('Preferred: ' + area if rank == 0 else area)}</td>"]
        for i, val in enumerate(values):
            if i == title_col and url and url != "#":
                cells.append(f"<td><a href='{_escape(url)}' target='_blank' rel='noopener'>{_escape(val)}</a></td>")
            elif i == url_col - 1:
                continue  # URL folded into the Job Title link, not shown as its own column
            else:
                cells.append(f"<td>{_escape(val)}</td>")
        cls = f" class='{row_class}'" if is_highlighted else ""
        rows_html_parts.append(f"<tr{cls}>{''.join(cells)}</tr>")

    rows_html = "\n".join(rows_html_parts) if rows_html_parts else f"<tr><td colspan='{len([c for c in columns if c != 'URL']) + 1}'>{_escape(empty_message)}</td></tr>"
    return rows_html, len(rows_html_parts)


def _section_html(title: str, columns: list[str], rows_html: str, total: int, legend_color: str, legend_label: str) -> str:
    display_columns = [c for c in columns if c != "URL"]
    header_html = "<th>Area priority</th>" + "".join(f"<th>{_escape(c)}</th>" for c in display_columns)
    return f"""
<h2>{_escape(title)}</h2>
<div class="meta">
  {total} tracked postings &middot;
  <span class="legend" style="background:{legend_color};"></span>{_escape(legend_label)}
</div>
<div class="scroll-wrap">
<table>
<thead><tr>{header_html}</tr></thead>
<tbody>
{rows_html}
</tbody>
</table>
</div>
"""



def _records(ws):
    """Presentation only: preserve the installed scope/seniority guards and ordering."""
    from src.matching import parse_posted
    if ws is None:return []
    records=[]
    profile=yaml.safe_load(CV_PROFILE_FILE.read_text(encoding="utf-8")) if CV_PROFILE_FILE.exists() else {}
    related_enabled=((profile or {}).get("related_roles") or {}).get("enabled") is True
    for row in range(2,ws.max_row+1):
        record={col:ws.cell(row,i+1).value for i,col in enumerate(COLUMNS)}
        if all(v is None for v in record.values()):continue
        if excluded_seniority(record.get('Job Title')) or not allowed_job({'city':record.get('City'),'company':record.get('Company'),'url':record.get('URL'),'location':record.get('Location')}):continue
        if record.get('Match Type')=='Related role' and not related_enabled:continue
        raw=record.get('Job Posted');posted=parse_posted(raw)
        score=record.get('Relevance Score (1-10)');score=score if isinstance(score,(int,float)) else 0
        if record.get('Match Type')=='Related role' and not 7 <= score <= 10:continue
        rank,area=area_priority(record.get('Location'))
        url=str(record.get('URL') or '')
        if urlparse(url).scheme not in ('http','https') or not urlparse(url).netloc:url=''
        records.append(dict(match_type=str(record.get('Match Type') or ''),match_reason=str(record.get('Match Reason') or ''),title=str(record.get('Job Title') or 'Untitled vacancy'),company=str(record.get('Company') or 'Employer unspecified'),location=str(record.get('Location') or 'Location not recorded'),score=score,rank=rank,area=area,url=url,posted=posted.isoformat() if posted else '',raw_date=str(raw or ''),new=_rgb_matches(ws.cell(row,1),NEW_ROW_FILL)))
    return sorted(records,key=lambda r:(r['rank'],-r['score']))

STYLE = r"""
:root{--bg:#f4f6f8;--ink:#172b36;--muted:#596d79;--line:#dce5e9;--accent:#086c65;--soft:#e7f4f0;--white:#fff}*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:Inter,-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;line-height:1.5}a{color:var(--accent)}button,input,select{font:inherit}button,select{cursor:pointer}button,a,input,select,summary{touch-action:manipulation}button:focus-visible,a:focus-visible,input:focus-visible,select:focus-visible,summary:focus-visible{outline:3px solid #c36b14;outline-offset:4px}[hidden]{display:none!important}.skip{position:absolute;left:16px;top:-80px;background:white;padding:12px;z-index:5}.skip:focus{top:12px}.topbar{background:#fff;border-bottom:1px solid var(--line)}.nav{max-width:1200px;margin:auto;padding:19px 28px;display:flex;align-items:center;justify-content:space-between;gap:20px}.brand{display:flex;align-items:center;gap:12px;font-weight:750;letter-spacing:-.4px}.mark{width:34px;height:34px;border-radius:11px;background:var(--accent);color:#fff;display:grid;place-items:center;font-size:18px}.nav small{color:var(--muted)}.nav a{font-size:13px;text-decoration:none}.shell{max-width:1200px;margin:auto;padding:42px 28px}.eyebrow{font-size:11px;letter-spacing:2px;font-weight:750;color:var(--accent);text-transform:uppercase}h1{font-size:clamp(30px,4vw,44px);letter-spacing:-1.6px;line-height:1.15;margin:10px 0 14px}h2{letter-spacing:-.5px}.hero{display:flex;align-items:flex-start;justify-content:space-between;gap:32px}.hero p{color:var(--muted);max-width:650px;margin:0 0 16px}.stamp{font-size:12px;color:var(--muted);padding:12px 15px;border:1px solid var(--line);border-radius:12px;background:white;white-space:nowrap}.stamp strong{display:block;color:var(--ink);font-weight:600}.scope{display:flex;flex-wrap:wrap;gap:8px}.scope span,.pill{display:inline-block;padding:5px 10px;border:1px solid var(--line);border-radius:30px;font-size:11px;background:white;color:var(--muted)}.metrics{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin:30px 0}.metric{background:white;border:1px solid var(--line);border-radius:16px;padding:19px 22px}.metric strong{display:block;font-size:29px;line-height:1.25;letter-spacing:-1px}.metric span{font-size:12px;color:var(--muted)}.workspace{display:grid;grid-template-columns:235px minmax(0,1fr);gap:26px;align-items:start}.filters{background:white;border:1px solid var(--line);border-radius:18px;padding:21px;position:sticky;top:20px}.filters h2{font-size:16px;margin:0 0 20px}.filters label{display:block;font-size:12px;font-weight:650;margin:18px 0 7px}.filters select,.filters input[type=search]{width:100%;border:1px solid #c9d7dc;border-radius:9px;background:#fff;color:var(--ink);padding:10px;min-height:43px}.filters .check{display:flex;align-items:center;gap:9px;font-weight:400;margin:18px 0}.filters input[type=checkbox]{accent-color:var(--accent);width:17px;height:17px}.hint{font-size:11px;color:var(--muted);line-height:1.6}.reset{border:0;background:transparent;color:var(--accent);padding:10px 0;font-size:12px;text-decoration:underline}.results-head{display:flex;align-items:center;justify-content:space-between;gap:16px;margin:0 0 16px}.results-head h2{font-size:19px;margin:0}.results-head p{font-size:12px;color:var(--muted);margin:3px 0}.results-head select{border:1px solid var(--line);background:white;padding:9px;border-radius:9px;font-size:12px;max-width:210px}.jobs{display:grid;gap:13px}.job-card{border:1px solid var(--line);border-radius:17px;background:white;padding:22px 24px;transition:border-color .15s,box-shadow .15s}.job-card:hover{border-color:#9bbdb7;box-shadow:0 6px 18px #172b3608}.job-top{display:flex;align-items:flex-start;justify-content:space-between;gap:16px}.employer{display:flex;align-items:center;gap:10px;font-size:12px;font-weight:650;color:var(--muted)}.avatar{display:grid;place-items:center;width:30px;height:30px;background:#f0f4f6;border:1px solid #e6ecef;border-radius:9px;font-size:10px;color:#526773}.job-card h3{font-size:19px;letter-spacing:-.35px;line-height:1.4;margin:12px 0 8px;overflow-wrap:anywhere}.job-card h3 a{color:var(--ink);text-decoration:none}.job-card h3 a:hover{text-decoration:underline}.score{background:var(--soft);border-radius:10px;padding:7px 11px;font-size:12px;white-space:nowrap;color:#075b55}.score b{font-size:17px}.location{font-size:12px;color:var(--muted);margin:0}.job-foot{display:flex;justify-content:space-between;align-items:center;gap:16px;margin-top:18px;padding-top:15px;border-top:1px solid #edf1f3}.badges{display:flex;align-items:center;flex-wrap:wrap;gap:7px}.pill{background:#f8fafb;border-color:#e3eaed}.pill.preferred{color:#075b55;background:#edf7f3;border-color:#cce6dc}.pill.new{color:#895207;background:#fff7e5;border-color:#f0dfb4}.apply{font-size:12px;text-decoration:none;font-weight:700;white-space:nowrap}.empty{background:white;border:1px dashed #b6c9ce;border-radius:16px;padding:40px;text-align:center;color:var(--muted)}.empty strong{display:block;color:var(--ink);margin-bottom:6px}.board-footer{font-size:11px;color:var(--muted);margin-top:22px}.health-panel{max-width:1144px;margin:0 auto 36px;border:1px solid var(--line);border-radius:18px;background:white;overflow:hidden}.health-panel>summary{padding:22px 24px;cursor:pointer;font-size:15px;font-weight:650}.health-panel>summary span{font-weight:400;color:var(--muted);font-size:12px;margin-left:10px}.health-content{padding:0 24px 24px}.health-content>p{font-size:12px;color:var(--muted)}.health-stats{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:15px}.health-stats span{font-size:11px;padding:6px 10px;background:#f3f6f8;border-radius:8px}.scroll-wrap{overflow:auto;max-height:520px;border:1px solid var(--line);border-radius:10px}table{border-collapse:collapse;width:100%;font-size:11px}th,td{text-align:left;vertical-align:top;padding:11px;border-bottom:1px solid var(--line)}th{background:#edf3f4;position:sticky;top:0;z-index:1;white-space:nowrap}td{min-width:100px}td:last-child{min-width:290px;max-width:550px;overflow-wrap:anywhere}.health-content details{font-size:11px}.health-content details summary{cursor:pointer;color:var(--accent)}.health-content details div{padding-top:10px}.noscript{padding:12px;background:#fff4da;font-size:12px}.demo-banner{background:#fff2d7;padding:10px 28px;font-size:12px;text-align:center}@media(max-width:800px){.shell{padding:28px 18px}.nav{padding:15px 18px}.hero{display:block}.stamp{display:inline-block;margin-top:18px}.metrics{grid-template-columns:repeat(2,1fr);gap:10px}.workspace{grid-template-columns:1fr}.filters{position:static}.filters h2{margin-bottom:12px}.filters label{margin-top:12px}.health-panel{margin:0 18px 28px}.results-head{align-items:flex-start}.job-card{padding:18px}.job-card h3{font-size:17px}.job-foot{align-items:flex-start}.nav small{display:none}.health-panel>summary span{display:block;margin:5px 0 0}}@media(prefers-reduced-motion:reduce){*{transition:none!important}}
"""


def _card(r,index):
    e=_escape
    title=f"<a href='{e(r['url'])}' target='_blank' rel='noopener noreferrer'>{e(r['title'])}</a>" if r['url'] else e(r['title'])
    link=f"<a class='apply' href='{e(r['url'])}' target='_blank' rel='noopener noreferrer' aria-label='View vacancy: {e(r['title'])}'>View vacancy &#8599;</a>" if r['url'] else "<span class='hint'>Link unavailable</span>"
    area={'Unterfoehring':'Unterföhring'}.get(r['area'],r['area'])
    date='Posted '+r['posted'] if r['posted'] else 'Posting date unknown'
    new="<span class='pill new'>New this update</span>" if r['new'] else ''
    related=("<span class='pill preferred' title='"+e(r.get('match_reason',''))+"'>Related role</span><span class='area-label'>"+e(r.get('match_reason',''))+"</span>") if r.get('match_type')=='Related role' else ''
    return f"""<article class='job-row'><div class='role'><h2>{title}</h2><p class='employer'>{e(r['company'])}</p>{new}{related}</div><div class='workplace'><span class='mobile-label'>Workplace</span>{e(r['location'])}<span class='area-label'>{e(area)}</span></div><div><span class='mobile-label'>Keyword relevance</span><span class='score' title='Existing weighted keyword score, not a hiring probability'><b>{r['score']:g}</b>/10</span></div><div class='posted'>{e(date)}</div><div class='action'>{link}</div></article>"""

LIST_STYLE=r"""
.shell{max-width:1340px;padding:25px 28px}.nav{max-width:1340px;padding:16px 28px}.list-heading{display:flex;align-items:flex-end;justify-content:space-between;gap:18px;margin-bottom:18px}.list-heading h1{font-size:27px;letter-spacing:-.8px;margin:0 0 5px}.list-heading p{margin:0;color:var(--muted);font-size:12px}.list-heading .stamp{font-size:11px;padding:9px 12px}.job-list{border:1px solid var(--line);border-radius:16px;background:white;overflow:hidden}.list-columns,.job-row{display:grid;grid-template-columns:minmax(0,2.6fr) minmax(0,1.3fr) 130px 150px 115px;gap:22px;align-items:center;padding:20px 24px}.list-columns{background:#edf3f4;padding-top:12px;padding-bottom:12px;font-size:10px;font-weight:700;text-transform:uppercase;letter-spacing:.7px;color:var(--muted)}.job-row{border-top:1px solid #e8eef1}.job-row:hover{background:#f7fbfa}.job-row h2{font-size:17px;line-height:1.45;letter-spacing:-.2px;margin:0 0 6px;overflow-wrap:anywhere}.job-row h2 a{color:var(--ink);text-decoration:none}.job-row h2 a:hover{text-decoration:underline}.job-row .employer{font-size:12px;margin:0 0 5px;font-weight:400}.workplace,.posted{font-size:12px;color:var(--muted);overflow-wrap:anywhere}.area-label{display:block;font-size:10px;color:var(--accent);margin-top:5px}.action{text-align:right}.mobile-label{display:none}.health-panel{max-width:1284px}.score{display:inline-block}.scope-note{font-size:11px;color:var(--muted);margin-top:14px}.empty{border:0}.job-row .pill.new{font-size:10px;padding:2px 8px}@media(max-width:1050px){.list-columns,.job-row{grid-template-columns:minmax(0,2fr) minmax(0,1fr) 100px 125px;gap:15px}.list-columns span:last-child{display:none}.action{grid-column:1/-1;text-align:left}}@media(max-width:700px){.shell{padding:22px 16px}.nav{padding:14px 16px}.list-heading{display:block}.list-heading .stamp{display:inline-block;margin-top:12px}.list-columns{display:none}.job-row{grid-template-columns:minmax(0,1fr) minmax(0,1fr);padding:20px 18px;gap:15px}.role{grid-column:1/-1}.job-row h2{font-size:18px}.mobile-label{display:block;font-size:10px;margin-bottom:6px;color:var(--muted)}.posted{grid-column:1/-1}.action{grid-column:1/-1}.health-panel{margin:0 16px 25px}.nav small{display:none}}
"""

def generate(tracker_path: Path = TRACKER_FILE, output_path: Path = OUTPUT_FILE) -> None:
    output_path=Path(output_path);output_path.parent.mkdir(parents=True,exist_ok=True)
    wb=None
    try:
        if Path(tracker_path).exists():wb=load_workbook(tracker_path)
        ws=wb['Jobs'] if wb and 'Jobs' in wb.sheetnames else (wb.active if wb else None)
        records=_records(ws)
    finally:
        if wb:wb.close()
    stamp=dt.datetime.now(dt.timezone.utc).strftime('%d %b %Y, %H:%M UTC')
    cards=''.join(_card(r,i) for i,r in enumerate(records)) or "<div class='empty'><strong>No tracked jobs to display</strong>Results appear after a successful tracker update.</div>"
    html=f"""<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Munich Job Tracker</title><style>{STYLE}{LIST_STYLE}</style></head><body><a class='skip' href='#jobs'>Skip to jobs</a><header class='topbar'><div class='nav'><div class='brand'><span class='mark' aria-hidden='true'>m.</span>Munich Job Tracker</div><small>Munich region · preferred areas first</small><a href='#scraper-health'>Source health &#8595;</a></div></header><main class='shell'><section aria-labelledby='jobs-title'><div class='list-heading'><div><h1 id='jobs-title'>Your opportunities</h1><p>{len(records)} tracked jobs · all shown · preferred areas, then relevance</p></div><div class='stamp'>Page generated<strong>{stamp}</strong>Not a guarantee of source freshness</div></div><div class='job-list' id='jobs'><div class='list-columns' aria-hidden='true'><span>Role &amp; employer</span><span>Workplace</span><span>Keyword relevance</span><span>Posting date</span><span>Vacancy</span></div>{cards}</div><p class='scope-note'>Munich metropolitan region · 21-day rule plus undated jobs · no commute cutoff. “New this update” means added in the latest tracker update, not newly posted. Keyword scores reflect existing rules, not hiring probability.</p></section></main></body></html>"""
    output_path.write_text(html,encoding='utf-8')

if __name__=='__main__':generate()
